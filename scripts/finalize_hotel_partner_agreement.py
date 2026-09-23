from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "public/legal/partner-agreements/Mandyal_Hotel_Partner_Agreement_v1.1.docx"
OUTPUT = ROOT / "public/legal/partner-agreements/Mandyal_Hotel_Partner_Agreement_v1.2.docx"


REPLACEMENTS = {
    "Draft for legal and operational review": "Management approved hotel partner agreement",
    "Review Copy 1.1 | 13 September 2026": "Version 1.2 | Effective 23 September 2026",
    (
        "This agreement is intended for onboarding an independent service provider to the Mandyal "
        "Travels platform. It becomes binding only when completed, duly executed, accepted by Mandyal "
        "Travels, and stamped or otherwise formalised as required by applicable law. The partner account "
        "remains inactive until written approval."
    ): (
        "This agreement governs onboarding of an independent hotel or accommodation provider to the "
        "Mandyal Travels platform. The Partner accepts this version through the platform's explicit "
        "electronic acceptance workflow or by returning a completed signed copy. The agreement becomes "
        "operative only when Mandyal Travels approves the application. Any stamp duty, signature, or "
        "other formality required for the particular transaction remains applicable. The partner account "
        "remains inactive until approval is recorded."
    ),
    (
        "[  ]  I will print or validly e-sign the emailed agreement, sign through an authorised "
        "representative, apply the business stamp if available, and return the complete copy from the "
        "registered email address."
    ): (
        "[  ]  I intend to accept this exact version electronically as the authorised representative. "
        "If Mandyal Travels identifies an applicable execution or stamping requirement, I will also "
        "return the requested complete signed or stamped copy from the registered business email."
    ),
    (
        "After submission, Mandyal Travels should send a locked, versioned copy of this Agreement from "
        "support@mandyaltravels.com to the verified partner email, showing the application reference, "
        "partner legal name and document version. The applicant should return the complete signed copy "
        "from that same verified email. A reply such as 'accepted' without the signed document is not "
        "sufficient for activation. Mandyal Travels should record the delivery timestamp, agreement hash "
        "or immutable version identifier, checkbox consent log, IP/device metadata where lawful, returned "
        "file, review notes and approval decision. The account remains pending until a named administrator "
        "records approval."
    ): (
        "Mandyal Travels will make a locked, versioned copy of this Agreement available to the applicant "
        "and may also send it to the verified partner email with the application reference, partner legal "
        "name, version and content hash. Electronic acceptance requires an unticked affirmative action by "
        "an authorised representative after the complete agreement is available for review. Mandyal "
        "Travels will retain the accepted name, user account, verified phone record, acceptance timestamp, "
        "agreement version and hash, and hashed request-device evidence where lawful. A generic email reply "
        "or a preselected checkbox is not acceptance. The account remains pending until a named "
        "administrator records approval."
    ),
    (
        "If the signed copy is not returned within 15 days, the application may expire and require fresh "
        "confirmation."
    ): (
        "If acceptance or requested supporting execution evidence is not completed within 15 days, the "
        "application may expire and require fresh confirmation."
    ),
    (
        "Mandyal Travels may require wet-ink signature, recognised electronic signature, notarisation or "
        "applicable e-stamping based on the transaction and state law."
    ): (
        "The platform acceptance record is evidence of contract acceptance; it is not represented as a "
        "licensed digital-signature certificate. Mandyal Travels may additionally require wet-ink or "
        "recognised electronic signature, notarisation, stamping or e-stamping where the instrument, "
        "transaction or applicable law requires that formality."
    ),
    (
        "The signatories confirm that they have authority, received the complete Agreement, had an "
        "opportunity to obtain independent advice, and intend to be legally bound when Mandyal Travels "
        "accepts the completed document. The parties will arrange applicable stamp duty or e-stamping as "
        "advised for the place and manner of execution."
    ): (
        "Each accepting representative confirms authority to bind the identified party, access to the "
        "complete Agreement, an opportunity to obtain independent advice, and an intention to be bound "
        "when Mandyal Travels records approval. Acceptance may be evidenced by the attributable electronic "
        "record described above or by signatures below. The parties remain responsible for any applicable "
        "stamp duty or execution formality for the place and manner of acceptance."
    ),
    (
        "This non-exhaustive reference list identifies the principal framework used for this review draft. "
        "It is not a representation that every listed law applies to every Partner or that every possible "
        "law has been listed. The Partner must identify and comply with all central, state, union-territory, "
        "local and mandatory foreign laws that apply to its actual service, location, route, licences, "
        "customers, data processing and operating model, including amendments. Locally qualified legal "
        "review is required before launching in a new state, union territory or foreign jurisdiction. "
        "International standards below are risk-management benchmarks unless incorporated into applicable "
        "law or a binding commitment."
    ): (
        "This non-exhaustive reference list identifies the principal framework considered for this "
        "Agreement. It is not a representation that every listed law applies to every Partner or that every "
        "possible law has been listed. The Partner must identify and comply with all central, state, "
        "union-territory, local and mandatory foreign laws that apply to its actual service, location, "
        "licences, customers, data processing and operating model, including amendments. Before launching "
        "in a new state, union territory or foreign jurisdiction, the parties must verify the local "
        "requirements and obtain appropriate professional advice where needed. International standards "
        "below are risk-management benchmarks unless incorporated into applicable law or a binding "
        "commitment."
    ),
}


def replace_paragraph(paragraph, replacement):
    if paragraph.text not in REPLACEMENTS:
        return False
    replacement = REPLACEMENTS[paragraph.text]
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)
    return True


document = Document(SOURCE)
matched = 0
for paragraph in document.paragraphs:
    matched += int(replace_paragraph(paragraph, REPLACEMENTS))
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                matched += int(replace_paragraph(paragraph, REPLACEMENTS))

if matched != len(REPLACEMENTS):
    raise RuntimeError(f"Expected {len(REPLACEMENTS)} replacements, applied {matched}")

document.core_properties.title = "Hotel and Accommodation Partner Services Agreement"
document.core_properties.subject = "Mandyal Travels hotel partner agreement version 1.2"
document.core_properties.comments = "Management-approved release for hotel partner onboarding"
document.save(OUTPUT)
print(OUTPUT)
